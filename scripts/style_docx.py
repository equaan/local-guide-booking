"""Post-processes the pandoc-generated .docx: fonts, headings, table borders, cover page, page numbers.

Usage: python scripts/style_docx.py path/to/report.docx
"""
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

NAVY = RGBColor(0x1F, 0x4E, 0x79)
BLUE = RGBColor(0x2F, 0x75, 0xA1)
DARK_GREY = RGBColor(0x33, 0x33, 0x33)
HEADER_BLUE = "1F4E79"
PALE_BLUE = "D9EAF7"
EVIDENCE_BLUE = "EAF3F8"
BODY_FONT = "Aptos"
CODE_FONT = "Consolas"


def set_font(style, name=None, size=None, bold=None, color=None):
    if name:
        style.font.name = name
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rfonts.set(qn(attr), name)
    if size:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color


def get_style(doc, name):
    try:
        return doc.styles[name]
    except KeyError:
        return None


def style_document(doc):
    for name in ("Normal", "Body Text", "First Paragraph", "Compact", "Block Text"):
        st = get_style(doc, name)
        if st is not None:
            set_font(st, BODY_FONT, 11)
            st.paragraph_format.space_after = Pt(6)
            st.paragraph_format.line_spacing = 1.15
    compact = get_style(doc, "Compact")
    if compact is not None:
        set_font(compact, BODY_FONT, 9)
        compact.paragraph_format.space_after = Pt(2)
        compact.paragraph_format.space_before = Pt(2)

    sizes = {"Heading 1": 17, "Heading 2": 13, "Heading 3": 11}
    for name, size in sizes.items():
        st = get_style(doc, name)
        if st is not None:
            set_font(st, BODY_FONT, size, bold=True, color=NAVY)
            st.paragraph_format.space_before = Pt(12 if name != "Heading 1" else 0)
            st.paragraph_format.space_after = Pt(6)
            st.paragraph_format.keep_with_next = True
    h1 = get_style(doc, "Heading 1")
    if h1 is not None:
        h1.paragraph_format.page_break_before = True

    for name in ("Source Code", "Verbatim Char"):
        st = get_style(doc, name)
        if st is not None:
            set_font(st, CODE_FONT, 9)

    cover = {
        "CoverTitle": (26, True, NAVY, 100),
        "CoverSub": (14, False, BLUE, 8),
        "CoverMeta": (14, True, NAVY, 52),
        "CoverSmall": (11, False, DARK_GREY, 8),
    }
    for name, (size, bold, color, before) in cover.items():
        st = get_style(doc, name)
        if st is not None:
            set_font(st, BODY_FONT, size, bold=bold, color=color)
            st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            st.paragraph_format.space_before = Pt(before)
            st.paragraph_format.space_after = Pt(6)
            st.paragraph_format.page_break_before = False

    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.left_margin = section.right_margin = Cm(2.0)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)


def shade(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def style_tables(doc):
    for table in doc.tables:
        tbl_pr = table._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), "808080")
            borders.append(el)
        tbl_pr.append(borders)

        cell_mar = OxmlElement("w:tblCellMar")
        for side, w in (("top", 40), ("left", 80), ("bottom", 40), ("right", 80)):
            el = OxmlElement(f"w:{side}")
            el.set(qn("w:w"), str(w))
            el.set(qn("w:type"), "dxa")
            cell_mar.append(el)
        tbl_pr.append(cell_mar)

        for r_idx, row in enumerate(table.rows):
            if r_idx == 0:
                tr_pr = row._tr.get_or_add_trPr()
                hdr = OxmlElement("w:tblHeader")
                tr_pr.append(hdr)
            for cell in row.cells:
                if r_idx == 0:
                    shade(cell, HEADER_BLUE)
                elif r_idx % 2 == 0:
                    shade(cell, PALE_BLUE)
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(9)
                        if r_idx == 0:
                            run.font.bold = True
                            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)


def style_screenshot_placeholders(doc):
    for paragraph in doc.paragraphs:
        if paragraph.style.name != "ScreenshotPlaceholder":
            continue
        p_pr = paragraph._p.get_or_add_pPr()
        shading = OxmlElement("w:shd")
        shading.set(qn("w:val"), "clear")
        shading.set(qn("w:color"), "auto")
        shading.set(qn("w:fill"), EVIDENCE_BLUE)
        p_pr.append(shading)
        border = OxmlElement("w:pBdr")
        for edge in ("top", "left", "bottom", "right"):
            element = OxmlElement(f"w:{edge}")
            element.set(qn("w:val"), "single")
            element.set(qn("w:sz"), "6")
            element.set(qn("w:space"), "6")
            element.set(qn("w:color"), "2F75A1")
            border.append(element)
        p_pr.append(border)
        paragraph.paragraph_format.space_before = Pt(8)
        paragraph.paragraph_format.space_after = Pt(8)
        paragraph.paragraph_format.line_spacing = 1.15


def add_page_numbers(doc):
    for section in doc.sections:
        footer = section.footer
        footer.is_linked_to_previous = False
        p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        for r in list(p.runs):
            r._r.getparent().remove(r._r)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        label = p.add_run("Local Guide Booking System | Page ")
        label.font.size = Pt(8.5)
        label.font.color.rgb = DARK_GREY
        run = p.add_run()
        run.font.size = Pt(8.5)
        run.font.color.rgb = DARK_GREY
        for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
            if kind:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), kind)
            else:
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = text
            run._r.append(el)


def scale_images(doc):
    max_width = Cm(17.0)
    for shape in doc.inline_shapes:
        if shape.width > max_width:
            ratio = max_width / shape.width
            shape.width = int(shape.width * ratio)
            shape.height = int(shape.height * ratio)


def main(path):
    doc = Document(path)
    style_document(doc)
    style_tables(doc)
    style_screenshot_placeholders(doc)
    scale_images(doc)
    add_page_numbers(doc)
    doc.save(path)


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit("usage: python scripts/style_docx.py report.docx")
    main(sys.argv[1])
